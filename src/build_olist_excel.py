"""
Build the Excel report for the REAL Olist data: excel/Olist_Real_Funnel_Report.xlsx

The two raw Kaggle tables are pasted as-is into the 'MQLs' and 'Deals' sheets.
The join between them is done IN EXCEL with MATCH/INDEX (no pre-joined data),
and every number on the report sheets is a formula.

Usage:  python src/build_olist_excel.py
"""

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.utils import get_column_letter

from build_excel_report import (BOX, DATE, F_BOLD, F_BODY, F_HEAD, F_INPUT, F_NOTE, FILL_HEAD,
                                FILL_INPUT, FILL_LIGHT, FONT, GREY, INT, NAVY, PCT, SUN, TEAL,
                                body, finish_chart, header, title, widths)

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "data" / "real_olist"
OUT = ROOT / "excel" / "Olist_Real_Funnel_Report.xlsx"
PP = '+0.0" pp";-0.0" pp";0.0" pp"'


def load():
    mql = pd.read_csv(SRC / "olist_marketing_qualified_leads_dataset.csv",
                      parse_dates=["first_contact_date"])
    deals = pd.read_csv(SRC / "olist_closed_deals_dataset.csv", parse_dates=["won_date"])
    mql["origin"] = mql["origin"].fillna("(missing)")
    return mql, deals


# --------------------------------------------------------------- data sheets --
def deals_sheet(wb, deals):
    ws = wb.create_sheet("Deals")
    cols = list(deals.columns)
    ws.append(cols)
    for row in deals.itertuples(index=False):
        ws.append([None if pd.isna(v) else (v.to_pydatetime() if isinstance(v, pd.Timestamp) else v)
                   for v in row])
    n = len(deals) + 1
    won_col = get_column_letter(cols.index("won_date") + 1)
    for r in range(2, n + 1):
        ws[f"{won_col}{r}"].number_format = "yyyy-mm-dd hh:mm"
    for c in ws[1]:
        c.font, c.fill = F_HEAD, FILL_HEAD
    for i in range(1, len(cols) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 18
    ws.freeze_panes = "A2"
    ws.sheet_properties.tabColor = GREY
    rng = lambda name: f"Deals!${get_column_letter(cols.index(name) + 1)}$2:${get_column_letter(cols.index(name) + 1)}${n}"
    return {c: rng(c) for c in cols}


def mql_sheet(wb, mql, D):
    """Raw MQLs + formula columns that join each lead to its deal."""
    ws = wb.create_sheet("MQLs")
    headers = ["mql_id", "first_contact_date", "lead_month", "landing_page_id", "origin",
               "deal_row", "is_won", "won_date", "sr_id", "days_to_close", "in_window"]
    ws.append(headers)
    for i, row in enumerate(mql.itertuples(index=False), start=2):
        ws.append([
            row.mql_id, row.first_contact_date.to_pydatetime(),
            row.first_contact_date.strftime("%Y-%m"), row.landing_page_id, row.origin,
            f"=IFERROR(MATCH(A{i},{D['mql_id']},0),\"\")",
            f"=IF(F{i}=\"\",0,1)",
            f"=IF(F{i}=\"\",\"\",INDEX({D['won_date']},F{i}))",
            f"=IF(F{i}=\"\",\"\",INDEX({D['sr_id']},F{i}))",
            f"=IF(F{i}=\"\",\"\",INT(H{i})-B{i})",
            f"=IF(B{i}>=Settings!$C$5,1,0)",
        ])
    n = len(mql) + 1
    for r in range(2, n + 1):
        ws[f"B{r}"].number_format = DATE
        ws[f"H{r}"].number_format = DATE
    for c in ws[1]:
        c.font, c.fill = F_HEAD, FILL_HEAD
    for col in "FGHIJK":
        ws[f"{col}1"].fill = PatternFill("solid", fgColor=TEAL)
    for i, w in enumerate([34, 16, 11, 34, 16, 9, 8, 12, 34, 13, 10], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.sheet_properties.tabColor = GREY
    return {h: f"MQLs!${get_column_letter(i)}$2:${get_column_letter(i)}${n}"
            for i, h in enumerate(headers, start=1)}


# ------------------------------------------------------------- report sheets --
def start_here(wb):
    ws = wb.active
    ws.title = "Start Here"
    title(ws, "Olist Marketing Funnel – Real Data Report",
          "Source: Marketing Funnel by Olist (Kaggle, CC BY-NC-SA 4.0) · analysis by Karan Lodha")
    rows = [
        ("Sheet", "What it answers"),
        ("Settings", "Analysis window start. Change the yellow cell only."),
        ("Monthly", "Lead volume and conversion by month – and why 2017 is excluded"),
        ("Channels", "Which marketing channels bring leads that become sellers?"),
        ("Landing Pages", "Do some sign-up pages produce much better leads?"),
        ("Page x Channel", "Is the page effect real, or just better traffic?"),
        ("Time to Close", "How long from first contact to a won deal?"),
        ("Closers", "How dependent are we on a few sales reps?"),
        ("Data Quality", "What we can and can't trust in this data"),
        ("MQLs / Deals", "Raw Kaggle tables. Teal columns on MQLs join each lead to its deal with MATCH/INDEX."),
    ]
    for i, (a, b) in enumerate(rows, start=4):
        if i == 4:
            header(ws, 4, [a, b])
            continue
        ws.cell(row=i, column=1, value=a)
        ws.cell(row=i, column=2, value=b)
        body(ws.cell(row=i, column=1), bold=True)
        body(ws.cell(row=i, column=2))
    ws.cell(row=16, column=1,
            value="Funnel: a marketing qualified lead (MQL) is a seller who filled in a sign-up "
                  "form; a deal is a seller the sales team signed onto the Olist marketplace.").font = F_NOTE
    widths(ws, [20, 100])


def settings(wb, M, D):
    ws = wb.create_sheet("Settings")
    title(ws, "Settings", "Change the yellow cell; every report sheet recalculates.")
    header(ws, 4, ["Setting", "", "Value", "Why"])
    ws.merge_cells("A4:B4")
    ws["A5"] = "Analysis window: leads from"
    ws["A5"].font = F_BODY
    ws.merge_cells("A5:B5")
    ws["C5"] = pd.Timestamp("2018-01-01").to_pydatetime()
    ws["C5"].font, ws["C5"].fill, ws["C5"].border, ws["C5"].number_format = F_INPUT, FILL_INPUT, BOX, DATE
    ws["D5"] = "Deal data starts Dec 2017, so earlier leads look like they never convert."
    ws["D5"].font = F_NOTE
    checks = [
        ("Leads in window", f"=SUM({M['in_window']})", INT),
        ("Deals in window", f"=COUNTIFS({M['in_window']},1,{M['is_won']},1)", INT),
        ("Conversion in window", "=C9/C8", PCT),
        ("Earliest deal in file", f"=MIN({D['won_date']})", DATE),
    ]
    for i, (label, f, fmt) in enumerate(checks, start=8):
        ws.cell(row=i, column=1, value=label).font = F_BODY
        c = ws.cell(row=i, column=3, value=f)
        c.number_format, c.font = fmt, F_BOLD
    widths(ws, [28, 6, 14, 70])


def monthly(wb, M, months):
    ws = wb.create_sheet("Monthly")
    title(ws, "Leads and Conversion by Month", "Conversion = leads that became a deal / leads")
    header(ws, 4, ["Lead month", "Leads", "Deals", "Conversion", "In window?"])
    for i, mth in enumerate(months):
        r = 5 + i
        ws.cell(row=r, column=1, value=mth)
        body(ws.cell(row=r, column=1), bold=True)
        cells = {
            2: (f"=COUNTIFS({M['lead_month']},$A{r})", INT),
            3: (f"=COUNTIFS({M['lead_month']},$A{r},{M['is_won']},1)", INT),
            4: (f"=IFERROR(C{r}/B{r},0)", PCT),
            5: (f"=IF(DATEVALUE($A{r}&\"-01\")>=Settings!$C$5,\"yes\",\"no – deal data incomplete\")", None),
        }
        for c, (f, fmt) in cells.items():
            ws.cell(row=r, column=c, value=f)
            body(ws.cell(row=r, column=c), fmt)
    last = 4 + len(months)
    ws.cell(row=last + 2, column=1,
            value="2017 conversion is near zero because the deals file starts in Dec 2017 – "
                  "a data-extract artefact, not a sales improvement in 2018.").font = F_NOTE
    ch = BarChart()
    ch.style = 10
    ch.title = "Conversion by lead month"
    ch.add_data(Reference(ws, min_col=4, min_row=4, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=last))
    ch.legend = None
    ch.y_axis.number_format = "0%"
    ch.series[0].graphicalProperties.solidFill = NAVY
    finish_chart(ch, labels=False)
    ch.height, ch.width = 8, 18
    ws.add_chart(ch, "G4")
    widths(ws, [12, 9, 9, 12, 26])


def channels(wb, M, origins):
    ws = wb.create_sheet("Channels")
    title(ws, "Channel Performance (analysis window)", "Origin = how the lead found the sign-up page")
    header(ws, 4, ["Channel (origin)", "Leads", "Share of leads", "Deals", "Conversion",
                   "vs average", "Avg days to close"])
    first = 5
    last = first + len(origins) - 1
    tot = last + 1
    w = M["in_window"]
    for i, o in enumerate(origins):
        r = first + i
        ws.cell(row=r, column=1, value=o)
        body(ws.cell(row=r, column=1), bold=True)
        cells = {
            2: (f"=COUNTIFS({M['origin']},$A{r},{w},1)", INT),
            3: (f"=IFERROR(B{r}/$B${tot},0)", PCT),
            4: (f"=COUNTIFS({M['origin']},$A{r},{w},1,{M['is_won']},1)", INT),
            5: (f"=IFERROR(D{r}/B{r},0)", PCT),
            6: (f"=(E{r}-$E${tot})*100", PP),
            7: (f"=IFERROR(AVERAGEIFS({M['days_to_close']},{M['origin']},$A{r},{w},1,"
                f"{M['days_to_close']},\">=0\"),0)", "0.0"),
        }
        for c, (f, fmt) in cells.items():
            ws.cell(row=r, column=c, value=f)
            body(ws.cell(row=r, column=c), fmt)
    ws.cell(row=tot, column=1, value="Total")
    for c, f, fmt in [(2, f"=SUM(B{first}:B{last})", INT), (3, f"=SUM(C{first}:C{last})", PCT),
                      (4, f"=SUM(D{first}:D{last})", INT), (5, f"=IFERROR(D{tot}/B{tot},0)", PCT)]:
        ws.cell(row=tot, column=c, value=f)
        body(ws.cell(row=tot, column=c), fmt, bold=True)
    body(ws.cell(row=tot, column=1), bold=True)
    for c in range(1, 8):
        ws.cell(row=tot, column=c).fill = FILL_LIGHT
    ws.conditional_formatting.add(f"F{first}:F{last}", ColorScaleRule(
        start_type="min", start_color="F8696B", mid_type="num", mid_value=0,
        mid_color="FFFFFF", end_type="max", end_color="63BE7B"))
    ws.cell(row=tot + 2, column=1,
            value="Read: social is 18% of leads but converts at less than half the average; "
                  "'unknown' converts best – a tracking gap that hides the best leads.").font = F_NOTE
    ch = BarChart()
    ch.type = "bar"
    ch.style = 10
    ch.title = "Lead → deal conversion by channel"
    ch.add_data(Reference(ws, min_col=5, min_row=4, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
    ch.legend = None
    ch.x_axis.scaling.orientation = "maxMin"
    ch.y_axis.number_format = "0%"
    ch.series[0].graphicalProperties.solidFill = TEAL
    finish_chart(ch)
    ch.dataLabels.numFmt = "0.0%"
    ch.height, ch.width = 9, 16
    ws.add_chart(ch, f"A{tot + 4}")
    widths(ws, [20, 9, 13, 9, 12, 12, 15])


def landing_pages(wb, M, pages):
    ws = wb.create_sheet("Landing Pages")
    title(ws, "Landing Page Performance (analysis window)",
          "The 10 pages with the most leads; the long tail is grouped below")
    header(ws, 4, ["Page", "Landing page id", "Leads", "Deals", "Conversion", "Share of leads"])
    w = M["in_window"]
    first, last = 5, 5 + len(pages) - 1
    other, tot = last + 1, last + 2
    for i, pid in enumerate(pages):
        r = first + i
        ws.cell(row=r, column=1, value=f"Page {i + 1}")
        ws.cell(row=r, column=2, value=pid)
        body(ws.cell(row=r, column=1), bold=True)
        body(ws.cell(row=r, column=2))
        ws.cell(row=r, column=2).font = Font(name=FONT, size=8, color=GREY)
        for c, f, fmt in [(3, f"=COUNTIFS({M['landing_page_id']},$B{r},{w},1)", INT),
                          (4, f"=COUNTIFS({M['landing_page_id']},$B{r},{w},1,{M['is_won']},1)", INT),
                          (5, f"=IFERROR(D{r}/C{r},0)", PCT),
                          (6, f"=IFERROR(C{r}/$C${tot},0)", PCT)]:
            ws.cell(row=r, column=c, value=f)
            body(ws.cell(row=r, column=c), fmt)
    ws.cell(row=other, column=1, value="All other pages")
    body(ws.cell(row=other, column=1), bold=True)
    for c, f, fmt in [(3, f"=C{tot}-SUM(C{first}:C{last})", INT),
                      (4, f"=D{tot}-SUM(D{first}:D{last})", INT),
                      (5, f"=IFERROR(D{other}/C{other},0)", PCT),
                      (6, f"=IFERROR(C{other}/$C${tot},0)", PCT)]:
        ws.cell(row=other, column=c, value=f)
        body(ws.cell(row=other, column=c), fmt)
    ws.cell(row=tot, column=1, value="Total")
    body(ws.cell(row=tot, column=1), bold=True)
    for c, f, fmt in [(3, f"=SUM({w})", INT),
                      (4, f"=COUNTIFS({w},1,{M['is_won']},1)", INT),
                      (5, f"=IFERROR(D{tot}/C{tot},0)", PCT)]:
        ws.cell(row=tot, column=c, value=f)
        body(ws.cell(row=tot, column=c), fmt, bold=True)
    for c in range(1, 7):
        ws.cell(row=tot, column=c).fill = FILL_LIGHT
    ws.conditional_formatting.add(f"E{first}:E{other}", ColorScaleRule(
        start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50,
        mid_color="FFEB84", end_type="max", end_color="63BE7B"))
    ws.cell(row=tot + 2, column=1,
            value="Pages 1 and 2 convert at ~22% vs ~10% for the rest. See 'Page x Channel' "
                  "for whether that is the page or the traffic.").font = F_NOTE
    widths(ws, [16, 36, 9, 9, 12, 14])


def page_by_channel(wb, M, top2, origins):
    ws = wb.create_sheet("Page x Channel")
    title(ws, "Top-2 Pages vs Other Pages, Within Each Channel",
          "If the top pages also win inside the same channel, the page itself is the driver")
    ws["A4"] = "Top page 1"
    ws["B4"] = top2[0]
    ws["A5"] = "Top page 2"
    ws["B5"] = top2[1]
    for r in (4, 5):
        ws[f"A{r}"].font = F_BOLD
        ws[f"B{r}"].font = Font(name=FONT, size=8, color=GREY)
    header(ws, 7, ["Channel", "Leads – top pages", "Conversion – top pages",
                   "Leads – other pages", "Conversion – other pages", "Difference"])
    w, p, o, won = M["in_window"], M["landing_page_id"], M["origin"], M["is_won"]
    for i, ch_name in enumerate(origins):
        r = 8 + i
        ws.cell(row=r, column=1, value=ch_name)
        body(ws.cell(row=r, column=1), bold=True)
        top_n = (f"COUNTIFS({o},$A{r},{w},1,{p},$B$4)+COUNTIFS({o},$A{r},{w},1,{p},$B$5)")
        top_w = (f"COUNTIFS({o},$A{r},{w},1,{p},$B$4,{won},1)"
                 f"+COUNTIFS({o},$A{r},{w},1,{p},$B$5,{won},1)")
        cells = {
            2: (f"={top_n}", INT),
            3: (f"=IFERROR(({top_w})/B{r},0)", PCT),
            4: (f"=COUNTIFS({o},$A{r},{w},1)-B{r}", INT),
            5: (f"=IFERROR((COUNTIFS({o},$A{r},{w},1,{won},1)-({top_w}))/D{r},0)", PCT),
            6: (f"=(C{r}-E{r})*100", PP),
        }
        for c, (f, fmt) in cells.items():
            ws.cell(row=r, column=c, value=f)
            body(ws.cell(row=r, column=c), fmt)
    last = 7 + len(origins)
    ws.conditional_formatting.add(f"F8:F{last}", DataBarRule(
        start_type="num", start_value=0, end_type="max", color="63BE7B"))
    ws.cell(row=last + 2, column=1,
            value="Top pages win by 10-15 points in organic, direct, referral and unknown traffic, "
                  "but not in paid search. Next step: A/B test sending organic traffic to the top page.").font = F_NOTE
    ch = BarChart()
    ch.style = 10
    ch.title = "Conversion: top pages vs other pages"
    ch.add_data(Reference(ws, min_col=3, min_row=7, max_row=last), titles_from_data=True)
    ch.add_data(Reference(ws, min_col=5, min_row=7, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=8, max_row=last))
    ch.y_axis.number_format = "0%"
    ch.series[0].graphicalProperties.solidFill = TEAL
    ch.series[1].graphicalProperties.solidFill = "B8C4CE"
    finish_chart(ch, labels=False)
    ch.height, ch.width = 8, 18
    ws.add_chart(ch, f"A{last + 4}")
    widths(ws, [18, 16, 20, 18, 22, 12])


def time_to_close(wb, M):
    ws = wb.create_sheet("Time to Close")
    title(ws, "Time to Close (analysis window)", "Days from first contact to won deal")
    header(ws, 4, ["Time to close", "From (days)", "To (days)", "Deals", "Share", "Cumulative"])
    buckets = [("0-7 days", 0, 7), ("8-30 days", 8, 30), ("31-90 days", 31, 90),
               ("Over 90 days", 91, 9999)]
    for i, (label, lo, hi) in enumerate(buckets):
        r = 5 + i
        ws.cell(row=r, column=1, value=label)
        body(ws.cell(row=r, column=1), bold=True)
        ws.cell(row=r, column=2, value=lo)
        ws.cell(row=r, column=3, value=hi)
        body(ws.cell(row=r, column=2), INT)
        body(ws.cell(row=r, column=3), INT)
        cells = {
            4: (f"=COUNTIFS({M['in_window']},1,{M['days_to_close']},\">=\"&B{r},"
                f"{M['days_to_close']},\"<=\"&C{r})", INT),
            5: (f"=IFERROR(D{r}/SUM($D$5:$D$8),0)", PCT),
            6: (f"=SUM($E$5:E{r})", PCT),
        }
        for c, (f, fmt) in cells.items():
            ws.cell(row=r, column=c, value=f)
            body(ws.cell(row=r, column=c), fmt)
    ws["A10"] = "Median days to close"
    ws["A10"].font = F_BOLD
    ws["D10"] = ArrayFormula("D10", f"=MEDIAN(IF(({M['in_window']}=1)*ISNUMBER({M['days_to_close']}),"
                                    f"{M['days_to_close']}))")
    ws["D10"].font = Font(name=FONT, size=12, bold=True, color=TEAL)
    ws["A11"] = "(array formula – in old Excel versions confirm with Ctrl+Shift+Enter)"
    ws["A11"].font = F_NOTE
    ch = BarChart()
    ch.style = 10
    ch.title = "Deals by time to close"
    ch.add_data(Reference(ws, min_col=4, min_row=4, max_row=8), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=8))
    ch.legend = None
    ch.series[0].graphicalProperties.solidFill = NAVY
    finish_chart(ch)
    ch.height, ch.width = 8, 14
    ws.add_chart(ch, "H4")
    widths(ws, [22, 12, 11, 9, 9, 12])


def closers(wb, M, reps):
    ws = wb.create_sheet("Closers")
    title(ws, "Sales Rep Concentration (analysis window)",
          "Deals closed per sales rep (sr_id), largest first – a Pareto view")
    header(ws, 4, ["Rank", "Sales rep id", "Deals", "Share", "Cumulative share"])
    first, last = 5, 5 + len(reps) - 1
    for i, rid in enumerate(reps):
        r = first + i
        ws.cell(row=r, column=1, value=i + 1)
        ws.cell(row=r, column=2, value=rid)
        body(ws.cell(row=r, column=1), INT, bold=True)
        body(ws.cell(row=r, column=2))
        ws.cell(row=r, column=2).font = Font(name=FONT, size=8, color=GREY)
        cells = {
            3: (f"=COUNTIFS({M['sr_id']},$B{r},{M['in_window']},1)", INT),
            4: (f"=IFERROR(C{r}/SUM($C${first}:$C${last}),0)", PCT),
            5: (f"=SUM($D${first}:D{r})", PCT),
        }
        for c, (f, fmt) in cells.items():
            ws.cell(row=r, column=c, value=f)
            body(ws.cell(row=r, column=c), fmt)
    ws.conditional_formatting.add(f"C{first}:C{last}", DataBarRule(
        start_type="num", start_value=0, end_type="max", color="5B9BD5"))
    ws["G5"] = "Reps needed for 50% of deals"
    ws["G5"].font = F_BOLD
    ws["J5"] = f"=COUNTIF(E{first}:E{last},\"<0.5\")+1"
    ws["J5"].font = Font(name=FONT, size=12, bold=True, color="C62828")
    ws["G6"] = "Total reps with deals"
    ws["G6"].font = F_BOLD
    ws["J6"] = f"=COUNTIF(C{first}:C{last},\">0\")"
    ws["J6"].font = F_BOLD
    ws["G8"] = "Key-person risk: losing the top rep would remove ~1 in 6 deals."
    ws["G8"].font = F_NOTE
    widths(ws, [7, 36, 9, 9, 15])


def data_quality(wb, M, D):
    ws = wb.create_sheet("Data Quality")
    title(ws, "Data Quality Checks", "Real data is never clean – these checks shape the conclusions")
    header(ws, 4, ["Check", "Value", "Share", "What it means"])
    n_m, n_d = "COUNTA(" + M["mql_id"] + ")", "COUNTA(" + D["mql_id"] + ")"
    rows = [
        ("Leads (MQLs)", f"={n_m}", None, "Jun 2017 – May 2018"),
        ("Deals", f"={n_d}", None, "Won Dec 2017 – Nov 2018: earlier deals are missing"),
        ("Deals that matched a lead", f"=SUM({M['is_won']})", f"=B7/B6",
         "Should be 100% – join integrity"),
        ("Leads with origin 'unknown' or missing",
         f"=COUNTIF({M['origin']},\"unknown\")+COUNTIF({M['origin']},\"(missing)\")", f"=B8/B5",
         "Tracking gap – and these leads convert best"),
        ("Deals won before first contact", f"=COUNTIF({M['days_to_close']},\"<0\")", f"=B9/B6",
         "Impossible date – excluded from timing"),
        ("Deals with declared revenue = 0", f"=COUNTIF({D['declared_monthly_revenue']},0)",
         f"=B10/B6", "Revenue field unusable – no value analysis"),
        ("Deals with has_company empty", f"=COUNTBLANK({D['has_company']})", f"=B11/B6",
         "Field mostly empty – not used"),
    ]
    for i, (label, f, share, note) in enumerate(rows, start=5):
        ws.cell(row=i, column=1, value=label)
        body(ws.cell(row=i, column=1), bold=True)
        ws.cell(row=i, column=2, value=f)
        body(ws.cell(row=i, column=2), INT)
        ws.cell(row=i, column=3, value=share)
        body(ws.cell(row=i, column=3), PCT)
        ws.cell(row=i, column=4, value=note)
        body(ws.cell(row=i, column=4))
    widths(ws, [38, 10, 9, 55])


def main():
    OUT.parent.mkdir(exist_ok=True)
    mql, deals = load()
    window = mql[mql["first_contact_date"] >= "2018-01-01"]
    won_ids = set(deals["mql_id"])
    months = sorted(mql["first_contact_date"].dt.strftime("%Y-%m").unique())
    origins = list(window["origin"].value_counts().index)
    pages = list(window["landing_page_id"].value_counts().head(10).index)
    top_channels = (window[window["landing_page_id"].isin(pages[:2])]["origin"]
                    .value_counts().loc[lambda s: s >= 50].index.tolist())
    reps = list(deals[deals["mql_id"].isin(window["mql_id"])]["sr_id"].value_counts().index)
    assert len(won_ids) == len(deals)

    wb = Workbook()
    start_here(wb)
    D = deals_sheet(wb, deals)
    M = mql_sheet(wb, mql, D)
    settings(wb, M, D)
    monthly(wb, M, months)
    channels(wb, M, origins)
    landing_pages(wb, M, pages)
    page_by_channel(wb, M, pages[:2], top_channels)
    time_to_close(wb, M)
    closers(wb, M, reps)
    data_quality(wb, M, D)

    report = ["Settings", "Monthly", "Channels", "Landing Pages", "Page x Channel",
              "Time to Close", "Closers", "Data Quality"]
    wb._sheets = [wb[n] for n in ["Start Here"] + report + ["MQLs", "Deals"]]
    wb.active = 0
    for n in report:
        wb[n].sheet_properties.tabColor = TEAL
    wb["Start Here"].sheet_properties.tabColor = SUN
    wb.calculation.fullCalcOnLoad = True   # Excel computes every formula on open
    wb.save(OUT)
    print(f"Saved {OUT}")


if __name__ == "__main__":
    main()
